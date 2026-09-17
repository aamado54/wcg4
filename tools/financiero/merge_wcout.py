"""Fusiona wcout2d nuevo con consolidado previo (períodos nuevos pisan al anterior)."""

from __future__ import annotations

import openpyxl
from pathlib import Path


def _row_key(row) -> tuple | None:
    if not row or len(row) < 3:
        return None
    bu, code = row[1], row[2]
    if bu is None or code is None:
        return None
    code = str(code).strip()
    if not code:
        return None
    newcode = row[4] if len(row) > 4 else 0
    return (str(bu).strip().upper(), code, newcode)


def _period_cols(headers: list) -> dict:
    return {
        h: i
        for i, h in enumerate(headers)
        if h and str(h).startswith("n") and str(h)[1:].isdigit()
    }


def _load_consolidado(path: Path) -> tuple[list, list, list]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["consolidado"]
    rows = list(ws.iter_rows(values_only=True))
    meta = []
    if "meta" in wb.sheetnames:
        meta = list(wb["meta"].iter_rows(values_only=True))
    wb.close()
    return list(rows[0]), rows[1:], meta


def merge_wcout(
    base_path: Path,
    new_path: Path,
    output_path: Path,
) -> list[str]:
    """Return merged period column names (nYYMM)."""
    old_h, old_d, _ = _load_consolidado(base_path)
    new_h, new_d, new_meta = _load_consolidado(new_path)

    old_p = _period_cols(old_h)
    new_p = _period_cols(new_h)
    override = set(new_p)
    merged_periods = sorted(
        [p for p in old_p if p not in override] + list(new_p),
        key=lambda x: int(x[1:]),
    )
    base_headers = list(old_h[:5]) + merged_periods

    old_by_key = {_row_key(r): list(r) for r in old_d if _row_key(r)}
    new_by_key = {_row_key(r): list(r) for r in new_d if _row_key(r)}

    def _val(row, pcols, period):
        if row and period in pcols and pcols[period] < len(row):
            return row[pcols[period]]
        return None

    def _build_row(key):
        ro, rn = old_by_key.get(key), new_by_key.get(key)
        base = rn or ro
        out = [None, base[1], base[2], base[3], base[4] if len(base) > 4 else 0]
        for p in merged_periods:
            if p in override:
                out.append(_val(rn, new_p, p))
            else:
                out.append(_val(ro, old_p, p))
        return out

    out_rows: list = []
    seen: set = set()
    for row in old_d:
        k = _row_key(row)
        if k is None:
            out_rows.append([None] * len(base_headers))
        elif k not in seen:
            out_rows.append(_build_row(k))
            seen.add(k)

    for k in new_by_key:
        if k not in seen and k[2] == 1:
            out_rows.append(_build_row(k))
            seen.add(k)

    wb_out = openpyxl.Workbook()
    ws = wb_out.active
    ws.title = "consolidado"
    ws.append(base_headers)
    for r in out_rows:
        ws.append(r)

    meta = wb_out.create_sheet("meta", 0)
    meta.append(["campo", "valor"])
    meta.append(["merged_from_base", str(base_path)])
    meta.append(["merged_from_new", str(new_path)])
    meta.append(["periods", ", ".join(merged_periods)])
    meta.append(["override_periods", ", ".join(sorted(override))])
    for mr in new_meta[1:]:
        meta.append(list(mr))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(output_path)
    return merged_periods
