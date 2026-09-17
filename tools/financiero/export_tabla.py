"""Exporta combined_full.json a una hoja Excel plana para revisión manual."""

from __future__ import annotations

import json
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter


def export_financiero_tabla(
    combined_full_path: Path,
    output_path: Path,
    *,
    include_kpis: bool = True,
) -> dict:
    data = json.loads(combined_full_path.read_text(encoding="utf-8"))
    accounts = data.get("accounts") or []
    periods = sorted(data.get("periods") or [])
    kpis = data.get("kpis") or {}

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EEFF consolidado"

    headers = ["BU", "Código", "Cuenta", "Fuente"] + periods
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for acc in sorted(accounts, key=lambda a: (a.get("bu", ""), a.get("code", ""))):
        vals = acc.get("values") or {}
        row = [
            acc.get("bu", ""),
            acc.get("code", ""),
            acc.get("label", ""),
            acc.get("source", ""),
        ]
        row.extend(vals.get(p) for p in periods)
        ws.append(row)

    if include_kpis:
        ws2 = wb.create_sheet("KPIs")
        kpi_names = sorted(
            {k for bu in kpis.values() for per in bu.values() for k in per}
        )
        ws2.append(["BU", "KPI"] + periods)
        for cell in ws2[1]:
            cell.font = Font(bold=True)
        for bu, bu_data in sorted(kpis.items()):
            for kpi in kpi_names:
                row = [bu, kpi]
                row.extend((bu_data.get(p) or {}).get(kpi) for p in periods)
                if any(v is not None for v in row[2:]):
                    ws2.append(row)

    meta = wb.create_sheet("meta", 0)
    meta.append(["campo", "valor"])
    meta.append(["source_json", str(combined_full_path)])
    meta.append(["generated", data.get("generated", "")])
    meta.append(["unit", data.get("unit", "")])
    meta.append(["accounts", len(accounts)])
    meta.append(["periods", ", ".join(periods)])
    meta.append(["recent_periods", ", ".join(data.get("recent_periods") or [])])

    for i, _ in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = 14
    ws.column_dimensions["C"].width = 42

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    return {
        "accounts": len(accounts),
        "periods": len(periods),
        "output": str(output_path),
    }
