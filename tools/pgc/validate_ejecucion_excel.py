#!/usr/bin/env python3
"""Cuadra PGC en DB vs archivo Excel de validación (hoja Objetivos).

Referencia mensual (columnas Excel en hoja Objetivos):
  S–V = puntos ejecución (ingresos, clientes, VC, reqs)
  W   = total puntos
  X   = Califica mes (1/0 si total ≥ 80) — **única calificación a validar**
  Y   = Califica trimestre — **ignorada** (no implementada en Tablero)

El Tablero en **2.Proporción** (`modo2`) debe cuadrar con S–W y X.
`modo1` (absolutos) no se compara con la hoja salvo con --mode modo1 explícito.
"""

from __future__ import annotations

import argparse
import os
import sys
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import django
import openpyxl

# Bootstrap Django when run as script
ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
os.environ.setdefault("DJANGO_SETTINGS_MODULE", "config.settings")
django.setup()

from core.models import MetricDefinition, UNE
from pgc.models import (
    MonthlyExchangeRate,
    MonthlyMetricResult,
    MonthlyMetricScore,
    MonthlyModeScorecard,
    MonthlyTarget,
    PGCPlan,
)

UEN_MAP = {
    "WC": "FACTORING",
    "WCL": "LEASING",
    "WCI": "INSURANCE",
    "WCINV": "INVESTMENT",
    "WCG": "TOTAL",
}

METRIC_MAP = {
    "clientes": MetricDefinition.CODE_CLIENTES_NUEVOS,
    "comisiones": MetricDefinition.CODE_INGRESOS,
}

# Índices 0-based en fila Objetivos (ver fila 1 del Excel Sep26)
COL_PTS_ING = 18   # S
COL_PTS_CLI = 19   # T
COL_PTS_VC = 20    # U
COL_PTS_REQ = 21   # V
COL_TOTAL = 22     # W
COL_CALIFICA_MES = 23  # X
# COL 24 (Y) Califica trimestre — no leer

SCORE_METRICS = (
    ("pts_ing", COL_PTS_ING, MetricDefinition.CODE_INGRESOS),
    ("pts_cli", COL_PTS_CLI, MetricDefinition.CODE_CLIENTES_NUEVOS),
    ("pts_vc", COL_PTS_VC, MetricDefinition.CODE_VENTA_CRUZADA),
    ("pts_req", COL_PTS_REQ, MetricDefinition.CODE_RESPUESTA_REQS),
)

POINTS_TOLERANCE = Decimal("0.05")
KNOWN_REQS_EXEMPT_UNES = frozenset({"INSURANCE"})


def _excel_bool_califica(raw) -> bool | None:
    if raw is None:
        return None
    if isinstance(raw, bool):
        return raw
    if isinstance(raw, (int, float, Decimal)):
        return Decimal(str(raw)) >= Decimal("1")
    s = str(raw).strip().upper()
    if s in ("1", "SI", "SÍ", "S", "TRUE", "CALIFICA", "YES"):
        return True
    if s in ("0", "NO", "N", "FALSE", "NO CALIFICA"):
        return False
    return None


def _dec_or_none(val):
    if val is None:
        return None
    try:
        return Decimal(str(val))
    except Exception:
        return None


def _read_objetivos(path: Path, year: int, month: int) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Objetivos"]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 10:
            continue
        dt = row[0]
        if not dt or not hasattr(dt, "year"):
            continue
        if dt.year != year or dt.month != month:
            continue
        uen = str(row[5]).strip().upper() if row[5] else ""
        rec = {
            "uen": uen,
            "pres_cli": row[1] if len(row) > 1 else None,
            "exec_cli": row[6] if len(row) > 6 else None,
            "pres_com_gtq": row[3] if len(row) > 3 else None,
            "exec_com_gtq": row[8] if len(row) > 8 else None,
        }
        if len(row) > COL_CALIFICA_MES:
            rec["pts_ing"] = row[COL_PTS_ING]
            rec["pts_cli"] = row[COL_PTS_CLI]
            rec["pts_vc"] = row[COL_PTS_VC]
            rec["pts_req"] = row[COL_PTS_REQ]
            rec["total_pts"] = row[COL_TOTAL]
            rec["califica_mes"] = row[COL_CALIFICA_MES]
        rows.append(rec)
    wb.close()
    return rows


def _gtq_comisiones_to_usd_miles(gtq: Decimal, fx: Decimal) -> Decimal:
    """Excel comisiones en GTQ completos → miles USD (como PGC)."""
    if not gtq or not fx:
        return Decimal("0")
    return (gtq / Decimal("1000")) / fx


def _db_value(plan, une_code: str, metric_code: str, year: int, month: int):
    une = UNE.objects.filter(code=une_code).first()
    metric = MetricDefinition.objects.filter(code=metric_code).first()
    if not une or not metric:
        return None, None
    res = MonthlyMetricResult.objects.filter(
        plan=plan, une=une, metric=metric, year=year, month=month
    ).first()
    tgt = MonthlyTarget.objects.filter(
        plan=plan, une=une, metric=metric, year=year, month=month
    ).first()
    return (
        res.measured_value if res else None,
        tgt.target_value if tgt else None,
    )


def _db_metric_points(plan, une_code: str, metric_code: str, year: int, month: int, mode: str):
    une = UNE.objects.filter(code=une_code).first()
    metric = MetricDefinition.objects.filter(code=metric_code).first()
    if not une or not metric:
        return None
    sc = MonthlyMetricScore.objects.filter(
        plan=plan,
        une=une,
        metric=metric,
        year=year,
        month=month,
        mode=mode,
    ).first()
    return sc.points_awarded if sc else None


def _db_scorecard(plan, une_code: str, year: int, month: int, mode: str):
    une = UNE.objects.filter(code=une_code).first()
    if not une:
        return None, None
    card = MonthlyModeScorecard.objects.filter(
        plan=plan, une=une, year=year, month=month, mode=mode
    ).first()
    if not card:
        return None, None
    return card.total_points, card.is_month_qualified


def compare_execution(path: Path, year: int, month: int) -> list[dict]:
    plan = PGCPlan.objects.filter(year=year).first()
    if not plan:
        raise SystemExit(f"No hay PGCPlan para {year}")

    fx_row = MonthlyExchangeRate.objects.filter(year=year, month=month).first()
    fx = fx_row.usd_to_gtq if fx_row else None

    findings = []
    for rec in _read_objetivos(path, year, month):
        uen = rec["uen"]
        if uen == "WCG":
            continue
        une_code = UEN_MAP.get(uen)
        if not une_code:
            findings.append({"uen": uen, "issue": "UEN no mapeada"})
            continue

        db_cli, tgt_cli = _db_value(
            plan, une_code, METRIC_MAP["clientes"], year, month
        )
        excel_cli = rec["exec_cli"]
        if excel_cli is not None and db_cli is not None:
            if Decimal(str(excel_cli)) != Decimal(str(db_cli)):
                findings.append(
                    {
                        "kind": "ejecucion",
                        "metric": "CLIENTES_NUEVOS",
                        "uen": uen,
                        "excel": excel_cli,
                        "db": db_cli,
                        "target": tgt_cli,
                        "excel_target": rec["pres_cli"],
                    }
                )

        if rec["exec_com_gtq"] is not None and une_code != "INVESTMENT":
            db_ing, tgt_ing = _db_value(
                plan, une_code, METRIC_MAP["comisiones"], year, month
            )
            excel_ing_usd = None
            if fx:
                excel_ing_usd = _gtq_comisiones_to_usd_miles(
                    Decimal(str(rec["exec_com_gtq"])), fx
                )
            if db_ing is None:
                findings.append(
                    {
                        "kind": "ejecucion",
                        "metric": "INGRESOS",
                        "uen": uen,
                        "excel_gtq": rec["exec_com_gtq"],
                        "excel_usd_miles": excel_ing_usd,
                        "db": None,
                        "target": tgt_ing,
                        "note": "Sin measured_value en DB",
                    }
                )
            elif excel_ing_usd is not None:
                diff = abs(Decimal(str(db_ing)) - excel_ing_usd)
                if diff > Decimal("0.05"):
                    findings.append(
                        {
                            "kind": "ejecucion",
                            "metric": "INGRESOS",
                            "uen": uen,
                            "excel_gtq": rec["exec_com_gtq"],
                            "excel_usd_miles": float(excel_ing_usd),
                            "db": float(db_ing),
                            "diff_usd": float(diff),
                            "target": float(tgt_ing) if tgt_ing else None,
                            "fx": float(fx),
                        }
                    )

    return findings


def compare_scores(
    path: Path, year: int, month: int, mode: str = "modo2"
) -> list[dict]:
    """Compara puntos S–W y calificación mes (columna X) vs DB."""
    plan = PGCPlan.objects.filter(year=year).first()
    if not plan:
        raise SystemExit(f"No hay PGCPlan para {year}")

    findings = []
    for rec in _read_objetivos(path, year, month):
        uen = rec["uen"]
        if uen in ("WCG", ""):
            continue
        une_code = UEN_MAP.get(uen)
        if not une_code:
            findings.append({"kind": "puntos", "uen": uen, "issue": "UEN no mapeada"})
            continue

        for field, _col, metric_code in SCORE_METRICS:
            excel_pts = _dec_or_none(rec.get(field))
            if excel_pts is None:
                continue
            db_pts = _db_metric_points(plan, une_code, metric_code, year, month, mode)
            if db_pts is None:
                findings.append(
                    {
                        "kind": "puntos",
                        "uen": uen,
                        "metric": metric_code,
                        "excel": float(excel_pts),
                        "db": None,
                        "note": f"Sin MonthlyMetricScore ({mode})",
                    }
                )
                continue
            diff = abs(Decimal(str(db_pts)) - excel_pts)
            if diff > POINTS_TOLERANCE:
                if (
                    metric_code == MetricDefinition.CODE_RESPUESTA_REQS
                    and une_code in KNOWN_REQS_EXEMPT_UNES
                    and db_pts > excel_pts
                ):
                    findings.append(
                        {
                            "kind": "puntos",
                            "uen": uen,
                            "metric": metric_code,
                            "excel": float(excel_pts),
                            "db": float(db_pts),
                            "diff": float(diff),
                            "note": "Excepción documentada Insurance reqs (app vs hoja)",
                            "severity": "info",
                        }
                    )
                else:
                    findings.append(
                        {
                            "kind": "puntos",
                            "uen": uen,
                            "metric": metric_code,
                            "excel": float(excel_pts),
                            "db": float(db_pts),
                            "diff": float(diff),
                        }
                    )

        excel_total = _dec_or_none(rec.get("total_pts"))
        db_total, db_califica = _db_scorecard(plan, une_code, year, month, mode)
        if excel_total is not None and db_total is not None:
            diff = abs(Decimal(str(db_total)) - excel_total)
            if diff > POINTS_TOLERANCE:
                findings.append(
                    {
                        "kind": "total",
                        "uen": uen,
                        "excel": float(excel_total),
                        "db": float(db_total),
                        "diff": float(diff),
                    }
                )
        elif excel_total is not None and db_total is None:
            findings.append(
                {
                    "kind": "total",
                    "uen": uen,
                    "excel": float(excel_total),
                    "db": None,
                    "note": f"Sin MonthlyModeScorecard ({mode})",
                }
            )

        excel_cal = _excel_bool_califica(rec.get("califica_mes"))
        if excel_cal is not None and db_califica is not None:
            if bool(db_califica) != excel_cal:
                findings.append(
                    {
                        "kind": "califica_mes",
                        "uen": uen,
                        "excel_col_x": rec.get("califica_mes"),
                        "excel_bool": excel_cal,
                        "db_is_month_qualified": db_califica,
                        "db_total": float(db_total) if db_total is not None else None,
                        "excel_total": float(excel_total) if excel_total is not None else None,
                    }
                )

    return findings


def compare(
    path: Path, year: int, month: int, mode: str = "modo2", include_execution: bool = True
) -> list[dict]:
    findings = compare_scores(path, year, month, mode=mode)
    if include_execution:
        findings.extend(compare_execution(path, year, month))
    return findings


def main():
    parser = argparse.ArgumentParser(
        description=(
            "Valida PGC vs Excel Objetivos. Calificación mensual: columna X (modo2). "
            "Columna Y (trimestre) no se usa."
        )
    )
    parser.add_argument("--path", required=True)
    parser.add_argument("--year", type=int, default=2026)
    parser.add_argument("--month", type=int, default=8)
    parser.add_argument(
        "--mode",
        default="modo2",
        choices=["modo1", "modo2"],
        help="Modalidad DB para puntos/califica (default modo2 = 2.Proporción, cuadra con Excel).",
    )
    parser.add_argument(
        "--scores-only",
        action="store_true",
        help="Solo puntos S–W y califica mes (X); omitir ejecución clientes/comisiones.",
    )
    parser.add_argument(
        "--months",
        type=str,
        default="",
        help="Lista de meses separados por coma (ej. 1,2,3) o 'all' para 1–12 presentes en Excel.",
    )
    args = parser.parse_args()

    path = Path(args.path)
    months: list[int]
    if args.months == "all":
        wb = openpyxl.load_workbook(path, data_only=True)
        ws = wb["Objetivos"]
        found = set()
        for row in ws.iter_rows(min_row=2, values_only=True):
            dt = row[0] if row else None
            if isinstance(dt, datetime) and dt.year == args.year:
                found.add(dt.month)
        wb.close()
        months = sorted(found) or [args.month]
    elif args.months:
        months = [int(m.strip()) for m in args.months.split(",") if m.strip()]
    else:
        months = [args.month]

    any_fail = False
    for month in months:
        findings = compare(
            path,
            args.year,
            month,
            mode=args.mode,
            include_execution=not args.scores_only,
        )
        # Errores reales (excluye severity info)
        hard = [f for f in findings if f.get("severity") != "info"]
        infos = [f for f in findings if f.get("severity") == "info"]

        if not findings:
            print(
                f"OK — {args.year}-{month:02d} ({args.mode}): "
                "puntos/califica mes (col X) y ejecución sin diferencias materiales."
            )
            continue

        if hard:
            any_fail = True
            print(
                f"Diferencias {args.year}-{month:02d} ({args.mode}) "
                f"— {len(hard)} hallazgo(s), col X = califica mes:"
            )
            for f in hard:
                print(" -", f)
        if infos:
            print(f"  Notas conocidas {args.year}-{month:02d} ({len(infos)}):")
            for f in infos:
                print(" -", f)

    return 1 if any_fail else 0


if __name__ == "__main__":
    raise SystemExit(main())
