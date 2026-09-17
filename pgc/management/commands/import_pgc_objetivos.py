"""Importa ejecución PGC (clientes + comisiones) desde hoja Objetivos de Excel."""

from __future__ import annotations

from datetime import datetime
from decimal import Decimal
from pathlib import Path

import openpyxl
from django.core.management.base import BaseCommand, CommandError
from django.db import transaction

from core.models import MetricDefinition, UNE
from pgc.income_conversion import apply_result_achievement, gtq_to_usd
from pgc.models import MonthlyExchangeRate, MonthlyMetricResult, MonthlyTarget, PGCPlan

UEN_TO_UNE = {
    "WC": "FACTORING",
    "WCL": "LEASING",
    "WCI": "INSURANCE",
    "WCINV": "INVESTMENT",
}


def _gtq_comisiones_to_usd_miles(gtq: Decimal, fx: Decimal) -> Decimal:
    """Comisiones en GTQ completos → miles USD (convención PGC tablero)."""
    if not gtq or not fx:
        return Decimal("0")
    return (gtq / Decimal("1000")) / fx


def _read_objetivos_rows(path: Path, year: int, month: int) -> list[dict]:
    wb = openpyxl.load_workbook(path, data_only=True)
    if "Objetivos" not in wb.sheetnames:
        wb.close()
        raise CommandError(f"{path.name}: falta hoja Objetivos")

    ws = wb["Objetivos"]
    rows: list[dict] = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or len(row) < 10:
            continue
        dt, _pres_cli, _pres_cli_ac, _pres_com, _pres_com_ac, uen, exec_cli, _exec_cli_ac, exec_com, _exec_com_ac = row[:10]
        if not isinstance(dt, datetime):
            continue
        if dt.year != year or dt.month != month:
            continue
        uen_code = str(uen or "").strip().upper()
        if uen_code not in UEN_TO_UNE:
            continue
        rows.append(
            {
                "uen": uen_code,
                "pres_cli": row[1] if len(row) > 1 else None,
                "pres_com_gtq": row[3] if len(row) > 3 else None,
                "exec_cli": exec_cli,
                "exec_com_gtq": exec_com,
            }
        )
    wb.close()
    return rows


class Command(BaseCommand):
    help = (
        "Carga ejecución mensual PGC (CLIENTES_NUEVOS e INGRESOS) desde la hoja "
        "Objetivos del Excel de validación / presupuesto."
    )

    def add_arguments(self, parser):
        parser.add_argument("--path", type=str, required=True, help="Ruta al .xlsx")
        parser.add_argument("--year", type=int, required=True)
        parser.add_argument("--month", type=int, required=True)
        parser.add_argument(
            "--sync-targets",
            action="store_true",
            help="Actualizar metas mensuales (presupuesto) desde columnas B y D del Excel.",
        )

    @transaction.atomic
    def handle(self, *args, **options):
        path = Path(options["path"]).expanduser()
        year = options["year"]
        month = options["month"]

        if not path.is_file():
            raise CommandError(f"Archivo no encontrado: {path}")

        try:
            plan = PGCPlan.objects.get(year=year)
        except PGCPlan.DoesNotExist:
            raise CommandError(f"No existe PGCPlan para {year}")

        fx_row = MonthlyExchangeRate.objects.filter(year=year, month=month).first()
        if not fx_row or not fx_row.usd_to_gtq:
            raise CommandError(
                f"Falta MonthlyExchangeRate para {year}-{month:02d}"
            )
        fx = fx_row.usd_to_gtq

        metric_cli = MetricDefinition.objects.get(code=MetricDefinition.CODE_CLIENTES_NUEVOS)
        metric_ing = MetricDefinition.objects.get(code=MetricDefinition.CODE_INGRESOS)

        records = _read_objetivos_rows(path, year, month)
        if not records:
            raise CommandError(f"Sin filas Objetivos para {year}-{month:02d}")

        sync_targets = options.get("sync_targets")

        for rec in records:
            une = UNE.objects.get(code=UEN_TO_UNE[rec["uen"]])

            if sync_targets:
                if rec.get("pres_cli") is not None:
                    tgt_cli = MonthlyTarget.objects.filter(
                        plan=plan, une=une, metric=metric_cli, year=year, month=month
                    ).first()
                    if tgt_cli:
                        tgt_cli.target_value = Decimal(str(rec["pres_cli"]))
                        tgt_cli.save(update_fields=["target_value", "updated_at"])
                if rec.get("pres_com_gtq") is not None and une.code != "INVESTMENT":
                    tgt_ing = MonthlyTarget.objects.filter(
                        plan=plan, une=une, metric=metric_ing, year=year, month=month
                    ).first()
                    if tgt_ing:
                        pres_usd = _gtq_comisiones_to_usd_miles(
                            Decimal(str(rec["pres_com_gtq"])), fx
                        )
                        tgt_ing.target_value = pres_usd
                        tgt_ing.save(update_fields=["target_value", "updated_at"])

            if rec["exec_cli"] is not None:
                tgt_cli = MonthlyTarget.objects.filter(
                    plan=plan, une=une, metric=metric_cli, year=year, month=month
                ).first()
                if not tgt_cli:
                    raise CommandError(
                        f"Sin meta CLIENTES_NUEVOS {une.code} {year}-{month:02d}"
                    )
                mmr_cli, _ = MonthlyMetricResult.objects.get_or_create(
                    plan=plan,
                    une=une,
                    metric=metric_cli,
                    year=year,
                    month=month,
                    defaults={"target_value": tgt_cli.target_value},
                )
                mmr_cli.target_value = tgt_cli.target_value
                mmr_cli.measured_value = Decimal(str(rec["exec_cli"]))
                mmr_cli.calculation_note = (
                    f"Importado desde {path.name} Objetivos "
                    f"(clientes ejecución mensual)."
                )
                apply_result_achievement(mmr_cli, tgt_cli)
                mmr_cli.save()

            if rec["exec_com_gtq"] is not None:
                tgt_ing = MonthlyTarget.objects.filter(
                    plan=plan, une=une, metric=metric_ing, year=year, month=month
                ).first()
                if not tgt_ing:
                    raise CommandError(
                        f"Sin meta INGRESOS {une.code} {year}-{month:02d}"
                    )
                gtq = Decimal(str(rec["exec_com_gtq"]))
                usd_miles = _gtq_comisiones_to_usd_miles(gtq, fx)
                mmr_ing, _ = MonthlyMetricResult.objects.get_or_create(
                    plan=plan,
                    une=une,
                    metric=metric_ing,
                    year=year,
                    month=month,
                    defaults={"target_value": tgt_ing.target_value},
                )
                mmr_ing.target_value = tgt_ing.target_value
                mmr_ing.source_currency = MonthlyMetricResult.CURRENCY_GTQ
                mmr_ing.source_value = gtq
                mmr_ing.exchange_rate_used = fx
                mmr_ing.conversion_status = MonthlyMetricResult.CONVERSION_CONVERTED
                mmr_ing.measured_value = usd_miles
                mmr_ing.calculation_note = (
                    f"Importado desde {path.name} Objetivos "
                    f"(comisiones ejecución mensual GTQ={gtq}, FX={fx}, "
                    f"USD miles={usd_miles})."
                )
                apply_result_achievement(mmr_ing, tgt_ing)
                mmr_ing.save()

            self.stdout.write(
                self.style.SUCCESS(
                    f"{rec['uen']} → clientes={rec['exec_cli']} "
                    f"comisiones_gtq={rec['exec_com_gtq']}"
                )
            )

        from django.core.management import call_command

        for mode in ("modo1", "modo2"):
            call_command("recalc_pgc", year=year, month=month, mode=mode)
        self.stdout.write(
            self.style.SUCCESS(f"Score PGC recalculado para {year}-{month:02d}.")
        )

        self.stdout.write(self.style.SUCCESS("Import PGC Objetivos completado."))
