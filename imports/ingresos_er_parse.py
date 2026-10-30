"""Lectura de comisiones mensuales (GTQ) desde ER operativos y PGA mensuales."""

from __future__ import annotations

import csv
import re
from decimal import Decimal, InvalidOperation
from pathlib import Path

import openpyxl
from django.core.management.base import CommandError

PGA_SHEET_BY_UNE = {
    "FACTORING": "FACTORING",
    "LEASING": "LEASING",
    "INSURANCE": "INSURANCE",
}

PGA_MONTH_ABBR = {
    1: "Ene",
    2: "Feb",
    3: "Mzo",
    4: "Abr",
    5: "May",
    6: "Jun",
    7: "Jul",
    8: "Ago",
    9: "Sep",
    10: "Oct",
    11: "Nov",
    12: "Dic",
}

MONTH_NAMES_ES = {
    1: "ENERO",
    2: "FEBRERO",
    3: "MARZO",
    4: "ABRIL",
    5: "MAYO",
    6: "JUNIO",
    7: "JULIO",
    8: "AGOSTO",
    9: "SEPTIEMBRE",
    10: "OCTUBRE",
    11: "NOVIEMBRE",
    12: "DICIEMBRE",
}

TSV_UNE_MAP = {
    "factoring": "FACTORING",
    "leasing": "LEASING",
    "insurance": "INSURANCE",
}
TSV_MONTHS = ("ene", "feb", "mzo", "abr", "may", "jun", "jul", "ago", "sep", "oct", "nov", "dic")
TSV_MONTH_NUM = {m: i + 1 for i, m in enumerate(TSV_MONTHS)}

from pgc.income_conversion import gtq_comisiones_to_usd_miles


def objetivos_style_ingresos_values(
    ingreso_mensual_gtq: Decimal, fx: Decimal
) -> tuple[Decimal, Decimal]:
    """
    Par source/measured alineado a import Objetivos y PGA.

    source_value = comisiones GTQ; measured_value = miles USD = (GTQ/1000)/FX.
    """
    if not fx or fx <= 0:
        raise ValueError("Tipo de cambio inválido para comisiones.")
    source_value = ingreso_mensual_gtq
    measured_value = gtq_comisiones_to_usd_miles(ingreso_mensual_gtq, fx)
    return source_value, measured_value


def is_pga_mensuales_workbook(path: Path) -> bool:
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        names = set(wb.sheetnames)
        wb.close()
    except Exception:
        return False
    return bool(names & set(PGA_SHEET_BY_UNE.values())) and "Datos" not in names


def _to_decimal(raw_value, context: str) -> Decimal:
    if raw_value in (None, ""):
        return Decimal("0")
    text = str(raw_value).strip().replace(",", "")
    try:
        return Decimal(text)
    except Exception as exc:
        raise CommandError(
            f"No se pudo convertir a decimal ({context}). Valor bruto={raw_value!r}"
        ) from exc


def _normalize_header(value) -> str:
    if value is None:
        return ""
    return str(value).strip().upper().replace(" ", "")


def _saldo_fin_col(ws) -> int | None:
    for col_idx, cell in enumerate(ws[1], start=1):
        header = _normalize_header(cell.value)
        if header in ("SALDOFIN", "SALDOFINAL"):
            return col_idx
    return None


def _month_col(ws, month: int) -> int | None:
    month_name = MONTH_NAMES_ES.get(month)
    if not month_name:
        return None
    for col_idx, cell in enumerate(ws[1], start=1):
        if cell.value is None:
            continue
        if str(cell.value).strip().upper() == month_name:
            return col_idx
    return None


def _numero_cuenta_col(ws) -> int:
    """Columna del código de cuenta (legacy B o NUMERO CUENTA)."""
    for col_idx, cell in enumerate(ws[1], start=1):
        header = _normalize_header(cell.value)
        if header in ("NUMEROCUENTA", "NROCUENTA", "CUENTA"):
            return col_idx
    return 2


def _tipo_col(ws) -> int:
    for col_idx, cell in enumerate(ws[1], start=1):
        if _normalize_header(cell.value) == "TIPO":
            return col_idx
    return 4


def _summary_saldo_for_prefix(ws, prefix: str) -> Decimal | None:
    """Fila resumen NUMERO CUENTA == 4 u 8 → SALDO FINAL."""
    saldo_fin_col = _saldo_fin_col(ws)
    if saldo_fin_col is None:
        return None
    nc_col = _numero_cuenta_col(ws)
    for row in ws.iter_rows(min_row=2):
        if nc_col - 1 >= len(row):
            continue
        cuenta = row[nc_col - 1].value
        if cuenta is None:
            continue
        if str(cuenta).strip() == prefix:
            saldo_value = row[saldo_fin_col - 1].value
            return _to_decimal(
                saldo_value,
                f"hoja Datos fila CUENTA={prefix} SALDO FINAL",
            )
    return None


def _sum_detail_month_column(ws, month: int, prefix: str) -> Decimal | None:
    """Suma cuentas detalle (9 dígitos) en columna mensual ENERO..DICIEMBRE."""
    month_col = _month_col(ws, month)
    if month_col is None:
        return None

    nc_col = _numero_cuenta_col(ws)
    tipo_col = _tipo_col(ws)
    total = Decimal("0")
    found = False
    for row in ws.iter_rows(min_row=2):
        if nc_col - 1 >= len(row):
            continue
        cuenta = row[nc_col - 1].value
        if cuenta is None:
            continue
        cuenta_str = str(cuenta).strip()
        cuenta_digits = "".join(ch for ch in cuenta_str if ch.isdigit())
        if not (cuenta_digits.startswith(prefix) and len(cuenta_digits) == 9):
            continue
        tipo = row[tipo_col - 1].value if tipo_col - 1 < len(row) else None
        if tipo != "D":
            continue
        month_value = row[month_col - 1].value
        if month_value in (None, ""):
            continue
        total += _to_decimal(
            month_value,
            f"cuenta {cuenta_str} columna mes {month}",
        )
        found = True
    return total if found else None


def _sum_detail_saldo_fin(ws, prefix: str) -> Decimal:
    """Suma detalle 9 dígitos en SALDO FINAL cuando no hay filas resumen."""
    saldo_fin_col = _saldo_fin_col(ws)
    if saldo_fin_col is None:
        return Decimal("0")
    nc_col = _numero_cuenta_col(ws)
    tipo_col = _tipo_col(ws)
    total = Decimal("0")
    for row in ws.iter_rows(min_row=2):
        if nc_col - 1 >= len(row):
            continue
        cuenta = row[nc_col - 1].value
        if cuenta is None:
            continue
        digits = "".join(ch for ch in str(cuenta).strip() if ch.isdigit())
        if len(digits) != 9 or not digits.startswith(prefix):
            continue
        tipo = row[tipo_col - 1].value if tipo_col - 1 < len(row) else None
        if tipo != "D":
            continue
        total += _to_decimal(
            row[saldo_fin_col - 1].value,
            f"detalle {cuenta}",
        )
    return total


def _saldo_fin_comisiones_total(ws) -> Decimal:
    """
    Comisiones del mes en ER WC* (cuentas 4+8).

    En archivos 2026+ SALDO FINAL ya es movimiento del mes, no YTD acumulado.
    """
    y4 = _summary_saldo_for_prefix(ws, "4")
    y8 = _summary_saldo_for_prefix(ws, "8")
    if y4 is not None or y8 is not None:
        return (y4 or Decimal("0")) + (y8 or Decimal("0"))
    return _sum_detail_saldo_fin(ws, "4") + _sum_detail_saldo_fin(ws, "8")


def infer_prev_er_path(path: Path, month: int) -> Path | None:
    if month <= 1:
        return None
    m = re.search(r"(\d{4})\.xlsx$", path.name, re.I)
    if not m:
        return None
    yymm = int(m.group(1))
    prev_yymm = yymm - 1
    if prev_yymm % 100 == 0:
        prev_yymm = (yymm // 100 - 1) * 100 + 12
    candidate = path.with_name(path.name.replace(f"{yymm:04d}", f"{prev_yymm:04d}"))
    return candidate if candidate.is_file() else None


def read_wc_er_comisiones_gtq(
    path: Path,
    month: int,
    *,
    prev_path: Path | None = None,
    monthly_saldo: bool = True,
) -> tuple[Decimal, str]:
    """
    Lee ingreso mensual GTQ (cuentas 4+8) desde hoja Datos de un ER WC*.

    Prioridad:
    1) Columna mensual ENERO..DICIEMBRE en cuentas detalle (legacy).
    2) SALDO FINAL filas resumen 4+8 (o detalle), como comisión del mes.
    3) Solo si monthly_saldo=False: delta YTD (actual − ER mes anterior).
    """
    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise CommandError(f"No se pudo abrir el archivo Excel: {e}")

    if "Datos" not in wb.sheetnames:
        wb.close()
        raise CommandError(
            f"El archivo {path.name} no contiene hoja 'Datos'. "
            f"Hojas disponibles: {', '.join(wb.sheetnames)}"
        )

    ws = wb["Datos"]
    monthly_4 = _sum_detail_month_column(ws, month, "4")
    monthly_8 = _sum_detail_month_column(ws, month, "8")
    if monthly_4 is not None or monthly_8 is not None:
        total = (monthly_4 or Decimal("0")) + (monthly_8 or Decimal("0"))
        wb.close()
        return total, "columna_mensual"

    saldo_total = _saldo_fin_comisiones_total(ws)
    wb.close()

    if monthly_saldo:
        return saldo_total, "saldo_fin_mensual"

    ytd_prev = Decimal("0")
    method = "ytd_sin_previo"
    if prev_path and prev_path.is_file():
        try:
            wb_prev = openpyxl.load_workbook(prev_path, data_only=True)
            if "Datos" in wb_prev.sheetnames:
                ytd_prev = _saldo_fin_comisiones_total(wb_prev["Datos"])
                method = "ytd_delta"
            wb_prev.close()
        except Exception:
            method = "ytd_sin_previo"

    return saldo_total - ytd_prev, method


def read_pga_mensuales_comisiones_gtq(
    path: Path, une_code: str, month: int
) -> Decimal:
    sheet = PGA_SHEET_BY_UNE.get(une_code.upper())
    if not sheet:
        raise CommandError(f"PGA mensuales no define hoja para UNE {une_code}.")
    abbr = PGA_MONTH_ABBR.get(month)
    if not abbr:
        raise CommandError(f"Mes inválido para PGA: {month}")

    try:
        wb = openpyxl.load_workbook(path, data_only=True)
    except Exception as e:
        raise CommandError(f"No se pudo abrir {path.name}: {e}")

    if sheet not in wb.sheetnames:
        wb.close()
        raise CommandError(f"{path.name}: falta hoja {sheet}.")

    ws = wb[sheet]
    month_col = None
    for row in ws.iter_rows(min_row=1, max_row=12, values_only=True):
        if not row:
            continue
        for idx, cell in enumerate(row):
            if cell is not None and str(cell).strip() == abbr:
                month_col = idx
                break
        if month_col is not None:
            break

    if month_col is None:
        wb.close()
        raise CommandError(
            f"{path.name} hoja {sheet}: no se encontró columna del mes {abbr}."
        )

    comisiones_row = None
    for row in ws.iter_rows(min_row=1, max_row=30, values_only=True):
        if not row or len(row) <= 2:
            continue
        label = row[2]
        if label is not None and str(label).strip().lower() == "comisiones":
            comisiones_row = row
            break

    wb.close()
    if comisiones_row is None:
        raise CommandError(f"{path.name} hoja {sheet}: falta fila «Comisiones».")

    if month_col >= len(comisiones_row):
        raise CommandError(f"{path.name}: columna {abbr} fuera de rango en Comisiones.")

    raw = comisiones_row[month_col]
    if raw in (None, ""):
        raise CommandError(
            f"{path.name} hoja {sheet}: sin comisión en {abbr}."
        )
    try:
        return Decimal(str(raw).replace(",", ""))
    except Exception:
        raise CommandError(
            f"{path.name}: comisión no numérica en {abbr}: {raw!r}"
        )


def _parse_gtq_cell(raw: str) -> Decimal | None:
    text = (raw or "").strip()
    if not text:
        return None
    text = text.replace(",", "").replace(" ", "")
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def load_comisiones_tsv(path: Path) -> dict[tuple[str, int, int], Decimal]:
    """(UNE_CODE, year, month) -> comisiones GTQ."""
    out: dict[tuple[str, int, int], Decimal] = {}
    year = 2026
    with path.open(newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f, delimiter="\t"))
    if len(rows) < 2:
        return out
    header = [c.strip().lower() for c in rows[0]]
    month_cols: list[tuple[int, int]] = []
    for idx, cell in enumerate(header):
        key = cell[:3] if len(cell) >= 3 else cell
        if key in TSV_MONTH_NUM:
            month_cols.append((idx, TSV_MONTH_NUM[key]))
    for row in rows[1:]:
        if not row or not row[0].strip():
            continue
        label = row[0].strip().lower()
        une = None
        for k, code in TSV_UNE_MAP.items():
            if k in label:
                une = code
                break
        if not une:
            continue
        for col_idx, month in month_cols:
            if col_idx >= len(row):
                continue
            gtq = _parse_gtq_cell(row[col_idx])
            if gtq is None:
                continue
            out[(une, year, month)] = gtq
    return out
