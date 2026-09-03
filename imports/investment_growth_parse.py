"""Lectura de archivos de crecimiento neto de inversiones (AP/PG) y préstamos bancarios."""

from __future__ import annotations

import csv
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string, get_column_letter

_PERIOD_RE = re.compile(r"^(20\d{2})[/-](0?[1-9]|1[0-2])$")
_FORMULA_RE = re.compile(
    r"=\(SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)\*([A-Z]+)(\d+)\)\+SUM\(([A-Z]+)(\d+):([A-Z]+)(\d+)\)",
    re.IGNORECASE,
)


def _parse_decimal(raw: str | None) -> Decimal | None:
    if raw is None:
        return None
    text = str(raw).strip().replace(" ", "").replace(",", "")
    if not text:
        return None
    try:
        return Decimal(text)
    except InvalidOperation:
        return None


def _parse_period(raw: str | None) -> tuple[int, int] | None:
    if not raw:
        return None
    text = str(raw).strip().replace("-", "/")
    match = _PERIOD_RE.match(text)
    if not match:
        return None
    return int(match.group(1)), int(match.group(2))


def _parse_date(raw: str | None) -> date | None:
    if not raw:
        return None
    text = str(raw).strip()
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%Y/%m/%d"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


@dataclass
class InvestmentGrowthRowData:
    year: int
    month: int
    instrument: str
    company: str
    operation_code: str
    currency_code: str
    amount_original: Decimal
    start_date: date | None
    maturity_date: date | None
    exchange_rate: Decimal | None
    amount_gtq: Decimal
    amount_usd: Decimal
    source_row_number: int | None = None


@dataclass
class BankMonthSnapshotData:
    year: int
    month: int
    exchange_rate: Decimal
    total_gtq: Decimal
    total_usd: Decimal
    bank_amounts: dict[str, Decimal] = field(default_factory=dict)
    usd_column_indices: tuple[int, ...] = ()
    gtq_column_indices: tuple[int, ...] = ()
    fx_column_index: int | None = None


def parse_inversiones_crecimiento_csv(path: Path) -> list[InvestmentGrowthRowData]:
    rows: list[InvestmentGrowthRowData] = []
    with path.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle, delimiter=";")
        header = next(reader, None)
        if not header:
            return rows

        for line_no, raw in enumerate(reader, start=2):
            if not raw or not any(str(c).strip() for c in raw):
                continue
            if len(raw) < 11:
                continue

            period = _parse_period(raw[0])
            if not period:
                continue
            year, month = period

            instrument = (raw[1] or "").strip().upper()
            if instrument not in ("AP", "PG"):
                continue

            amount_original = _parse_decimal(raw[5]) or Decimal("0")
            amount_gtq = _parse_decimal(raw[9]) or Decimal("0")
            amount_usd = _parse_decimal(raw[10]) or Decimal("0")
            if amount_gtq == 0 and amount_usd == 0:
                continue

            rows.append(
                InvestmentGrowthRowData(
                    year=year,
                    month=month,
                    instrument=instrument,
                    company=(raw[2] or "").strip(),
                    operation_code=(raw[3] or "").strip(),
                    currency_code=(raw[4] or "").strip().upper(),
                    amount_original=amount_original,
                    start_date=_parse_date(raw[6]),
                    maturity_date=_parse_date(raw[7]),
                    exchange_rate=_parse_decimal(raw[8]),
                    amount_gtq=amount_gtq,
                    amount_usd=amount_usd,
                    source_row_number=line_no,
                )
            )
    return rows


def _col_index(letter: str) -> int:
    return column_index_from_string(letter.upper())


def _parse_quetzalizado_formula(formula: str, row_number: int) -> tuple[tuple[int, ...], tuple[int, ...], int]:
    match = _FORMULA_RE.search(str(formula or ""))
    if not match:
        raise ValueError(f"No se pudo interpretar fórmula Quetzalizado: {formula}")

    usd_start = _col_index(match.group(1))
    usd_end = _col_index(match.group(3))
    fx_col = _col_index(match.group(5))
    gtq_start = _col_index(match.group(7))
    gtq_end = _col_index(match.group(9))

    if match.group(2) != match.group(4) or match.group(8) != match.group(10):
        raise ValueError(f"Rangos de fórmula inconsistentes en fila {row_number}")

    usd_cols = tuple(range(usd_start, usd_end + 1))
    gtq_cols = tuple(range(gtq_start, gtq_end + 1))
    return usd_cols, gtq_cols, fx_col


def _find_quetzalizado_column(ws) -> int | None:
    for row in ws.iter_rows(min_row=1, max_row=3):
        for cell in row:
            if cell.value and "quetzalizado" in str(cell.value).strip().lower():
                return cell.column
    return None


def parse_bancos_fin_mes_xlsx(path: Path) -> list[BankMonthSnapshotData]:
    wb_formula = load_workbook(path, data_only=False)
    wb_values = load_workbook(path, data_only=True)
    ws_formula = wb_formula.active
    ws_values = wb_values.active

    qtz_col = _find_quetzalizado_column(ws_formula)
    if not qtz_col:
        raise ValueError("No se encontró columna 'Quetzalizado' en Bancos_Fin_de_mes.")

    header_row = 1
    for row_idx in range(1, 4):
        if ws_formula.cell(row_idx, qtz_col).value and "quetzalizado" in str(
            ws_formula.cell(row_idx, qtz_col).value
        ).lower():
            header_row = row_idx
            break

    sample_formula = None
    sample_row = header_row + 1
    for row_idx in range(header_row + 1, header_row + 6):
        value = ws_formula.cell(row_idx, qtz_col).value
        if isinstance(value, str) and value.startswith("="):
            sample_formula = value
            sample_row = row_idx
            break
    if not sample_formula:
        raise ValueError("No se encontró fórmula Quetzalizado para detectar rangos USD/GTQ.")

    usd_cols, gtq_cols, fx_col = _parse_quetzalizado_formula(sample_formula, sample_row)

    bank_names: dict[int, str] = {}
    for col_idx in range(2, qtz_col):
        label = ws_values.cell(header_row, col_idx).value
        if label:
            bank_names[col_idx] = str(label).strip()

    snapshots: list[BankMonthSnapshotData] = []
    for row_idx in range(header_row + 1, ws_values.max_row + 1):
        period_raw = ws_values.cell(row_idx, 1).value
        period = _parse_period(str(period_raw) if period_raw is not None else None)
        if not period:
            continue

        year, month = period
        fx = _parse_decimal(ws_values.cell(row_idx, fx_col).value)
        if not fx or fx == 0:
            continue

        bank_amounts: dict[str, Decimal] = {}
        usd_sum = Decimal("0")
        gtq_sum = Decimal("0")
        for col_idx, name in bank_names.items():
            amount = _parse_decimal(ws_values.cell(row_idx, col_idx).value) or Decimal("0")
            if amount == 0:
                continue
            bank_amounts[name] = amount
            if col_idx in usd_cols:
                usd_sum += amount
            elif col_idx in gtq_cols:
                gtq_sum += amount

        total_gtq = ws_values.cell(row_idx, qtz_col).value
        if total_gtq is None:
            total_gtq = usd_sum * fx + gtq_sum
        total_gtq = Decimal(str(total_gtq))
        total_usd = total_gtq / fx

        snapshots.append(
            BankMonthSnapshotData(
                year=year,
                month=month,
                exchange_rate=fx,
                total_gtq=total_gtq,
                total_usd=total_usd,
                bank_amounts=bank_amounts,
                usd_column_indices=usd_cols,
                gtq_column_indices=gtq_cols,
                fx_column_index=fx_col,
            )
        )

    return snapshots


def verify_instrument_totals_gtq(
    rows: list[InvestmentGrowthRowData],
    expected: dict[tuple[int, int], dict[str, Decimal]],
) -> list[str]:
    """Compara totales AP/PG (quetzalizados) contra hoja de verificación."""
    totals: dict[tuple[int, int], dict[str, Decimal]] = {}
    for row in rows:
        bucket = totals.setdefault((row.year, row.month), {"AP": Decimal("0"), "PG": Decimal("0")})
        bucket[row.instrument] = bucket.get(row.instrument, Decimal("0")) + row.amount_gtq

    messages: list[str] = []
    for period, expected_vals in expected.items():
        actual = totals.get(period, {})
        for instrument in ("AP", "PG"):
            exp = expected_vals.get(instrument)
            if exp is None:
                continue
            got = actual.get(instrument, Decimal("0"))
            diff = abs(got - exp)
            if diff > Decimal("0.05"):
                messages.append(
                    f"{period[0]}-{period[1]:02d} {instrument}: archivo={got} verificación={exp} Δ={diff}"
                )
    return messages
