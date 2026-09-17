"""Utilidades compartidas para leer plantillas Excel (ocultos, hojas, celdas)."""

from __future__ import annotations

import re
from pathlib import Path

from openpyxl.utils import get_column_letter


def is_hidden(ws, row: int | None = None, col: int | None = None) -> bool:
    if ws.sheet_state == "hidden":
        return True
    if row and ws.row_dimensions[row].hidden:
        return True
    if col:
        letter = get_column_letter(col)
        if ws.column_dimensions[letter].hidden:
            return True
    return False


def safe_cell(ws, row: int, col: int):
    try:
        return ws.cell(row=row, column=col)
    except Exception:
        return None


def find_sheet(wb, names: list[str]):
    for ws in wb.worksheets:
        if ws.sheet_state == "hidden":
            continue
        norm = (
            ws.title.lower()
            .replace("á", "a")
            .replace("é", "e")
            .replace("í", "i")
            .replace("ó", "o")
            .replace("ú", "u")
        )
        for name in names:
            if norm.startswith(name):
                return ws
    return None


def unique_code(filename: str, used: set[str]) -> str:
    base = Path(filename).stem
    code = re.sub(r"\W+", "", base)[:8].upper() or "CLIENTE"
    orig = code
    i = 1
    while code in used:
        code = f"{orig[:7]}{i}"
        i += 1
    used.add(code)
    return code


def clean_company_name(filename: str) -> str:
    name = re.sub(r"\.(xlsx|xlsm|xls)$", "", filename.strip(), flags=re.I)
    name = re.sub(r"\s+v\d+\s*$", "", name, flags=re.I).strip()
    name = re.sub(r"\s+PFJuridica\s*$", "", name, flags=re.I).strip()
    # LU3-American Medical Enterprises → American Medical Enterprises
    m = re.match(r"^[A-Z]{2,4}\d*[-_\s]+(.+)$", name)
    if m:
        return m.group(1).strip()
    return name


def column_has_numbers(ws, col: int, start_row: int, end_row: int, min_numbers: int = 3) -> bool:
    count = 0
    for row in range(start_row, end_row + 1):
        if is_hidden(ws, row=row, col=col):
            continue
        cell = safe_cell(ws, row, col)
        if cell is None or cell.value is None:
            continue
        val = cell.value
        if isinstance(val, (int, float)):
            count += 1
        elif isinstance(val, str):
            try:
                float(val.replace(",", "."))
                count += 1
            except ValueError:
                pass
    return count >= min_numbers


def column_is_percentage(ws, col: int, start_row: int, end_row: int) -> bool:
    total = percent = 0
    for row in range(start_row, end_row + 1):
        if is_hidden(ws, row=row, col=col):
            continue
        cell = safe_cell(ws, row, col)
        if cell is None or cell.value is None:
            continue
        total += 1
        if isinstance(cell.value, str) and "%" in cell.value:
            percent += 1
    return total > 0 and (percent / total) > 0.8


def identify_version(ws) -> int | None:
    for col in range(1, 4):
        for row in range(1, 81):
            cell = safe_cell(ws, row, col)
            if cell and cell.value and isinstance(cell.value, str):
                val = cell.value.lower()
                if any(x in val for x in ("nombre del cliente", "ventas", "activo corriente")):
                    return 1
    for col in range(7, 10):
        for row in range(1, 81):
            cell = safe_cell(ws, row, col)
            if cell and cell.value and isinstance(cell.value, str):
                val = cell.value.lower()
                if any(x in val for x in ("nombre del cliente", "ventas", "activo corriente")):
                    return 2
    return None
