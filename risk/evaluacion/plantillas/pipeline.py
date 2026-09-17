"""Pipeline completo: directorio de plantillas → workbook evaluación WCG."""

from __future__ import annotations

import shutil
from dataclasses import dataclass, field
from pathlib import Path

import openpyxl
from openpyxl import Workbook

from .extract import TemplateBlock, extract_directory
from .transform import build_wcg_caratula_sheet, write_sheet_from_grid, write_side_by_side


@dataclass
class ConsolidateResult:
    source_dir: str
    output_path: str
    files_processed: int
    companies_in_caratula: int
    alerts: list[str] = field(default_factory=list)
    sheets: list[str] = field(default_factory=list)


def _copy_altman_from_reference(wb_out: Workbook, reference: Path | None) -> None:
    if reference and reference.is_file():
        ref = openpyxl.load_workbook(reference, data_only=True)
        if "Altman" in ref.sheetnames:
            ws_src = ref["Altman"]
            ws_dst = wb_out.create_sheet("Altman")
            for row in ws_src.iter_rows(values_only=True):
                ws_dst.append(list(row))
        ref.close()
        return
    # Fallback mínimo
    ws = wb_out.create_sheet("Altman")
    ws.append(["", "", "", ""])
    ws.append(["", "Tabla Resumida de Rangos de Interpretación", "", ""])
    ws.append(["", "", "", "Bajo Riesgo", "Riesgo Moderado", "Alto Riesgo"])


def build_evaluacion_workbook(
    blocks: list[TemplateBlock],
    output_path: Path,
    *,
    reference_workbook: Path | None = None,
    include_intermediate: bool = True,
) -> ConsolidateResult:
    wb_out = Workbook()
    wb_out.remove(wb_out.active)

    # Caratula (formato wcg4)
    caratula_grid = build_wcg_caratula_sheet(blocks)
    build_alerts: list[str] = []
    if caratula_grid:
        ws_c = wb_out.create_sheet("Caratula")
        write_sheet_from_grid(ws_c, caratula_grid)
    else:
        build_alerts.append("No se pudo construir hoja Caratula")

    _copy_altman_from_reference(wb_out, reference_workbook)

    if include_intermediate:
        v1: dict[str, list] = {}
        v2: dict[str, list] = {}
        bal_v2: dict[str, list] = {}
        res_v2: dict[str, list] = {}
        for block in blocks:
            if not block.caratula:
                continue
            if block.version == 1:
                v1[block.filename] = block.caratula
            else:
                v2[block.filename] = block.caratula
            if block.balance:
                bal_v2[block.filename] = block.balance
            if block.resultados:
                res_v2[block.filename] = block.resultados

        if v1:
            ws = wb_out.create_sheet("Caratula_V1")
            write_side_by_side(ws, v1, desc_cols=1)
        if v2:
            ws = wb_out.create_sheet("Caratula_V2")
            write_side_by_side(ws, v2, desc_cols=1)
        if bal_v2:
            ws = wb_out.create_sheet("Balance")
            write_side_by_side(ws, bal_v2, desc_cols=2)
        if res_v2:
            ws = wb_out.create_sheet("Resultados")
            write_side_by_side(ws, res_v2, desc_cols=2)


    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb_out.save(output_path)
    wb_out.close()

    companies = 0
    if caratula_grid and caratula_grid[0]:
        companies = sum(
            1 for v in caratula_grid[0] if isinstance(v, str) and v.strip()
        )

    return ConsolidateResult(
        source_dir="",
        output_path=str(output_path),
        files_processed=len(blocks),
        companies_in_caratula=companies,
        alerts=build_alerts,
        sheets=[s for s in openpyxl.load_workbook(output_path).sheetnames],
    )


def consolidate_templates(
    source_dir: Path,
    output_path: Path,
    *,
    reference_workbook: Path | None = None,
) -> ConsolidateResult:
    blocks, alerts = extract_directory(source_dir)
    result = build_evaluacion_workbook(
        blocks,
        output_path,
        reference_workbook=reference_workbook,
    )
    result.source_dir = str(source_dir)
    result.alerts = alerts
    if result.alerts:
        wb = openpyxl.load_workbook(output_path)
        if "Alertas" not in wb.sheetnames:
            ws_a = wb.create_sheet("Alertas")
            for a in result.alerts:
                ws_a.append([a])
            wb.save(output_path)
        wb.close()
    return result
